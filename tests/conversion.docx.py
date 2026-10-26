"""Independent development-only DOCX acceptance reader; not an app dependency."""

import json
import io
import pathlib
import sys
import zipfile
import xml.etree.ElementTree as ET

import docx

path = pathlib.Path(sys.argv[1])
expected = json.loads(pathlib.Path(sys.argv[2]).read_text(encoding="utf-8"))
word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
with zipfile.ZipFile(path) as package:
    assert package.testzip() is None, "DOCX ZIP CRC failure"
    names = set(package.namelist())
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names
    assert not any(name.startswith("/") or ".." in pathlib.PurePosixPath(name).parts for name in names)
    types = ET.fromstring(package.read("[Content_Types].xml"))
    assert any(part.attrib.get("PartName") == "/word/document.xml"
               and part.attrib.get("ContentType") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
               for part in types)
    relationships = ET.fromstring(package.read("_rels/.rels"))
    assert any(part.attrib.get("Type", "").endswith("/officeDocument")
               and part.attrib.get("Target") == "word/document.xml" for part in relationships)
    assert not any(part.attrib.get("TargetMode") == "External" for part in relationships)
    document = ET.fromstring(package.read("word/document.xml"))
    assert document.tag == word + "document"
    assert not document.findall(".//" + word + "drawing")
    assert not document.findall(".//" + word + "altChunk")
    breaks = document.findall(".//" + word + "br")
    page_breaks = sum(item.attrib.get(word + "type") == "page" for item in breaks)

opened = docx.Document(path)
pages = [[]]
for paragraph in opened.paragraphs:
    if any(item.attrib.get(word + "type") == "page" for item in paragraph._p.findall(".//" + word + "br")):
        pages.append([])
    else:
        pages[-1].append(paragraph.text)
page_text = ["\n".join(paragraphs) for paragraphs in pages]
assert len(page_text) == len(expected["pages"]), "Wrong selected-page count/order"
assert page_breaks == len(expected["pages"]) - 1
for index, (actual, original) in enumerate(zip(page_text, expected["pages"])):
    normalized = original.replace("\r\n", "\n").replace("\r", "\n")
    assert actual == normalized, f"Extracted Unicode text changed on selected page {index + 1}"
combined = "\n".join(page_text)
for fragment in expected.get("requiredFragments", []):
    assert fragment in combined, f"Missing Unicode/XML literal: {fragment}"
for fragment in expected.get("forbiddenFragments", []):
    assert fragment not in combined, f"Exported unselected/unsupported content: {fragment}"
opened.paragraphs[0].add_run("Independent editable round-trip marker 91827")
edited = io.BytesIO()
opened.save(edited)
edited.seek(0)
assert "Independent editable round-trip marker 91827" in docx.Document(edited).paragraphs[0].text
report = {
    "file": str(path),
    "reader": "python-docx " + docx.__version__,
    "zipCrcAndOpcXml": True,
    "exactEditableUnicodePageText": True,
    "independentTextEditSaveReopen": True,
    "pages": len(page_text),
    "pageBreaks": page_breaks,
    "paragraphs": len(opened.paragraphs),
    "characters": len(combined),
}
print(json.dumps(report, ensure_ascii=False))
